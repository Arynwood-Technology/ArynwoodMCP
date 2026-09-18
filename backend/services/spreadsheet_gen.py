"""
Deterministic .xlsx generation for the generate_spreadsheet native tool (chat.py).

Deliberately NOT "execute whatever Python the model wrote" — that would mean running
LLM-generated code server-side, a real security problem. Instead the model calls this
tool with structured data (title/headers/rows), and this fixed, tested code does all
the actual styling — the Arynwood-branded palette, number formatting, and the total-row
positioning bug (an off-by-one from computing the total row's index via len(data)
instead of the sheet's real row count, confirmed live while hand-tuning Glyph's prompt
before this tool existed) are all handled once here instead of regenerated — and
possibly re-broken — by the model on every call.
"""
import os
import re
import uuid

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from backend._frozen import user_data_dir

# Arynwood's real brand palette (frontend/src/index.css's @theme block: --color-accent
# is a plum/violet, --color-accent2 is a teal) — deepened here for fill-color contrast
# with white header text; the raw --color-accent2 (#5eead4) is a pastel, too light to
# read against white. "Plum teal blue" per the user's own framing: blue sits as the
# natural mid-tone on a plum-to-teal spectrum, used for secondary emphasis.
PLUM = "6B46C1"
BLUE = "2563EB"
TEAL = "0D9488"
LIGHT_BAND = "F5F3FF"      # very light plum tint, for alternating row shading
TOTAL_FILL = "CCFBF1"      # light teal tint, for the total row
TEXT_DARK = "1F1B2E"
WHITE = "FFFFFF"

_OUTPUT_DIR = os.path.join(user_data_dir(), "generated", "spreadsheets")
_MAX_ROWS = 5000  # sanity cap — this tool is for reports, not database dumps


class SpreadsheetError(ValueError):
    pass


def _safe_filename(title: str) -> str:
    slug = re.sub(r"[^a-zA-Z0-9]+", "_", title).strip("_").lower() or "spreadsheet"
    return f"{slug[:60]}_{uuid.uuid4().hex[:8]}.xlsx"


def _is_number(value) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def build_spreadsheet(
    title: str,
    headers: list[str],
    rows: list[list],
    number_columns: list[str] | None = None,
    currency_columns: list[str] | None = None,
    total_row: bool = False,
) -> tuple[str, str]:
    """Builds a styled .xlsx file. Returns (filename, absolute_path)."""
    if not headers:
        raise SpreadsheetError("headers must be a non-empty list")
    if not rows:
        raise SpreadsheetError("rows must be a non-empty list")
    if len(rows) > _MAX_ROWS:
        raise SpreadsheetError(f"too many rows ({len(rows)}) — this tool is for reports, not raw dumps (cap {_MAX_ROWS})")
    for i, row in enumerate(rows):
        if len(row) != len(headers):
            raise SpreadsheetError(f"row {i} has {len(row)} values, expected {len(headers)} (one per header)")

    number_columns = set(number_columns or [])
    currency_columns = set(currency_columns or [])
    unknown = (number_columns | currency_columns) - set(headers)
    if unknown:
        raise SpreadsheetError(f"number_columns/currency_columns reference headers that don't exist: {sorted(unknown)}")

    wb = Workbook()
    ws = wb.active
    ws.title = (title or "Sheet1")[:31]  # Excel's own sheet-name length limit

    header_font = Font(bold=True, color=WHITE)
    header_fill = PatternFill(start_color=PLUM, end_color=PLUM, fill_type="solid")
    thin = Side(style="thin", color="D8D3EC")
    cell_border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col_idx, name in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=name)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = cell_border
        cell.alignment = Alignment(horizontal="center", vertical="center")

    band_fill = PatternFill(start_color=LIGHT_BAND, end_color=LIGHT_BAND, fill_type="solid")
    for r_idx, row in enumerate(rows, start=2):
        banded = (r_idx % 2 == 0)
        for c_idx, (name, value) in enumerate(zip(headers, row), start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            cell.border = cell_border
            cell.font = Font(color=TEXT_DARK)
            if banded:
                cell.fill = band_fill
            if name in currency_columns and _is_number(value):
                cell.number_format = '"$"#,##0.00'
            elif name in number_columns and _is_number(value):
                cell.number_format = "#,##0"

    if total_row:
        # Computed from ws.max_row *after* every data row is already written, not
        # from len(rows) — the two only coincide when nothing else shares the sheet,
        # and silently diverge (an off-by-one) the moment that stops being true.
        total_r = ws.max_row + 1
        total_fill = PatternFill(start_color=TOTAL_FILL, end_color=TOTAL_FILL, fill_type="solid")
        top_border = Border(left=thin, right=thin, top=Side(style="medium", color=BLUE), bottom=thin)
        label_written = False
        for c_idx, name in enumerate(headers, start=1):
            is_numeric_col = name in currency_columns or name in number_columns
            if is_numeric_col:
                total = sum(
                    row[c_idx - 1] for row in rows
                    if _is_number(row[c_idx - 1])
                )
                cell = ws.cell(row=total_r, column=c_idx, value=total)
                cell.number_format = '"$"#,##0.00' if name in currency_columns else "#,##0"
            elif not label_written:
                cell = ws.cell(row=total_r, column=c_idx, value="Total")
                label_written = True
            else:
                cell = ws.cell(row=total_r, column=c_idx, value=None)
            cell.font = Font(bold=True, color=TEXT_DARK)
            cell.fill = total_fill
            cell.border = top_border

    for col_idx, name in enumerate(headers, start=1):
        values = [str(name)] + [str(row[col_idx - 1]) for row in rows]
        width = max(len(v) for v in values) + 4
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(width, 10), 40)

    ws.freeze_panes = "A2"

    os.makedirs(_OUTPUT_DIR, exist_ok=True)
    filename = _safe_filename(title)
    path = os.path.join(_OUTPUT_DIR, filename)
    wb.save(path)
    return filename, path


def resolve_download_path(filename: str) -> str:
    """Resolves a previously-generated filename to a real path, sandboxed to the
    generated-spreadsheets directory — filename is UUID-suffixed server-side (see
    _safe_filename) so this never trusts a client-supplied path, only a basename
    lookup within a directory this module itself controls."""
    if os.path.basename(filename) != filename:
        raise SpreadsheetError("invalid filename")
    path = os.path.join(_OUTPUT_DIR, filename)
    if not os.path.isfile(path):
        raise SpreadsheetError("not found")
    return path
