import openpyxl
import pytest

from backend.services import spreadsheet_gen as sg


@pytest.fixture(autouse=True)
def _isolated_output_dir(tmp_path, monkeypatch):
    """Redirect generated files into a throwaway dir so tests never touch (or
    depend on) the real generated/spreadsheets/ directory."""
    monkeypatch.setattr(sg, "_OUTPUT_DIR", str(tmp_path))


def _load(path):
    wb = openpyxl.load_workbook(path)
    return wb.active


def test_basic_generation_roundtrips_values():
    filename, path = sg.build_spreadsheet(
        title="Test Report",
        headers=["Name", "Score"],
        rows=[["Alice", 10], ["Bob", 20]],
    )
    assert filename.endswith(".xlsx")
    ws = _load(path)
    assert ws.title == "Test Report"
    assert [c.value for c in ws[1]] == ["Name", "Score"]
    assert [c.value for c in ws[2]] == ["Alice", 10]
    assert [c.value for c in ws[3]] == ["Bob", 20]


def test_header_row_styled():
    _, path = sg.build_spreadsheet(title="X", headers=["A"], rows=[[1]])
    ws = _load(path)
    header = ws["A1"]
    assert header.font.bold is True
    assert header.fill.fill_type == "solid"
    assert header.fill.start_color.rgb == f"00{sg.PLUM}"


def test_total_row_lands_on_correct_row_not_off_by_one():
    # The exact bug class this tool exists to eliminate: an LLM regenerating this
    # logic every call kept computing the total row's index from len(rows) instead
    # of the sheet's real row count, landing the total on the last data row instead
    # of its own row. Fixed code path here computes it from ws.max_row after every
    # data row is already written, so this must hold regardless of row count.
    for n in (1, 2, 5, 10):
        rows = [[f"item{i}", i * 10] for i in range(n)]
        _, path = sg.build_spreadsheet(
            title="X", headers=["Item", "Value"], rows=rows,
            number_columns=["Value"], total_row=True,
        )
        ws = _load(path)
        expected_total_row = n + 2  # header (1) + n data rows, total is the next row
        assert ws.cell(row=expected_total_row, column=1).value == "Total"
        assert ws.cell(row=expected_total_row, column=2).value == sum(i * 10 for i in range(n))
        # nothing left over below the total row
        assert ws.cell(row=expected_total_row + 1, column=1).value is None


def test_number_formatting_applied_consistently_including_total():
    _, path = sg.build_spreadsheet(
        title="X", headers=["Item", "Amount"],
        rows=[["a", 100], ["b", 200], ["c", 300]],
        currency_columns=["Amount"], total_row=True,
    )
    ws = _load(path)
    # every data cell and the total cell share the same currency format — the
    # exact inconsistency ("total reads $600.00, rows above read 100 unformatted")
    # this tool exists to make structurally impossible.
    amount_cells = [ws.cell(row=r, column=2) for r in range(2, 6)]
    formats = {c.number_format for c in amount_cells}
    assert formats == {'"$"#,##0.00'}


def test_currency_vs_plain_number_formatting_differ():
    _, path = sg.build_spreadsheet(
        title="X", headers=["Count", "Price"],
        rows=[[5, 9.99]],
        number_columns=["Count"], currency_columns=["Price"],
    )
    ws = _load(path)
    assert ws["A2"].number_format == "#,##0"
    assert ws["B2"].number_format == '"$"#,##0.00'


def test_freeze_panes_and_column_widths():
    _, path = sg.build_spreadsheet(
        title="X", headers=["Short", "A Much Longer Header Name"],
        rows=[["a", "b"]],
    )
    ws = _load(path)
    assert ws.freeze_panes == "A2"
    assert ws.column_dimensions["B"].width > ws.column_dimensions["A"].width


def test_empty_headers_rejected():
    with pytest.raises(sg.SpreadsheetError):
        sg.build_spreadsheet(title="X", headers=[], rows=[[1]])


def test_empty_rows_rejected():
    with pytest.raises(sg.SpreadsheetError):
        sg.build_spreadsheet(title="X", headers=["A"], rows=[])


def test_mismatched_row_length_rejected():
    with pytest.raises(sg.SpreadsheetError, match="row 1"):
        sg.build_spreadsheet(title="X", headers=["A", "B"], rows=[["ok", 1], ["too", "many", "values"]])


def test_unknown_column_name_in_number_columns_rejected():
    with pytest.raises(sg.SpreadsheetError, match="don't exist"):
        sg.build_spreadsheet(title="X", headers=["A"], rows=[[1]], number_columns=["NotAHeader"])


def test_too_many_rows_rejected():
    rows = [[i] for i in range(sg._MAX_ROWS + 1)]
    with pytest.raises(sg.SpreadsheetError, match="too many rows"):
        sg.build_spreadsheet(title="X", headers=["N"], rows=rows)


def test_resolve_download_path_roundtrip():
    filename, path = sg.build_spreadsheet(title="X", headers=["A"], rows=[[1]])
    resolved = sg.resolve_download_path(filename)
    assert resolved == path


def test_resolve_download_path_rejects_traversal():
    with pytest.raises(sg.SpreadsheetError):
        sg.resolve_download_path("../../etc/passwd")


def test_resolve_download_path_rejects_missing_file():
    with pytest.raises(sg.SpreadsheetError):
        sg.resolve_download_path("does_not_exist.xlsx")
